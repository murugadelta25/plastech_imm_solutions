from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import extract, func
from datetime import date, datetime, timedelta
from typing import Any, Dict, Optional, List
from pydantic import BaseModel, Field
import io, json, smtplib, os
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from ..models import ProductionPlan, Machine, ModelChangeRequest, EmailLog, ToolStock, get_db, now_ist, WorkOrder
from ..auth import get_current_user, require_role
from ..ws_manager import manager
from .work_orders import sync_work_order_after_plan_change

router = APIRouter(prefix="/api/plans", tags=["plans"])

class PlanCreate(BaseModel):
    plan_date: date
    end_date: Optional[date] = None   # if set, creates plans for each day in range
    shift: str
    shifts: Optional[List[str]] = None  # if set, creates one plan per day per shift
    station_no: int
    machine_id: Optional[int] = None
    work_order_id: Optional[int] = None
    mold_id: Optional[int] = None  # IMM mold master (tool_stocks.id)
    mold_details: Optional[Dict[str, Any]] = None  # snapshot of mold planning fields
    # Legacy CNC fields — optional for IMM (auto-filled from mold)
    current_operation: Optional[str] = None
    next_operation: Optional[str] = None
    model_variant: Optional[str] = None
    process_time: Optional[float] = None
    loading_unloading: float = 10
    planned_qty: int
    priority: int = 1
    plan_type: str = "scheduled"
    notes: Optional[str] = None
    tool_shortage_ack: bool = False  # planner ack when forecast requires it

class PlanUpdate(BaseModel):
    planned_qty: Optional[int] = None
    priority: Optional[int] = None
    plan_type: Optional[str] = None
    notes: Optional[str] = None
    status: Optional[str] = None

class ActualQtyUpdate(BaseModel):
    actual_qty: int
    source: str = "manual"  # manual | modbus | opcua | mqtt

class RescheduleRequest(BaseModel):
    new_date: Optional[date] = None
    new_shift: Optional[str] = None
    new_machine_id: Optional[int] = None
    new_station_no: Optional[int] = None
    mode: str = "custom"  # custom | next_week | split_remaining
    days_offset: Optional[int] = None
    split_remaining: bool = False

class BulkRescheduleRequest(BaseModel):
    plan_ids: List[int]
    mode: str = "next_week"  # next_week | custom
    days_offset: int = 7
    new_start_date: Optional[date] = None
    split_remaining: bool = True

def _resolve_shifts(data: PlanCreate) -> List[str]:
    if data.shifts:
        return list(dict.fromkeys(data.shifts))
    return [data.shift]

def _plan_copy_fields(plan: ProductionPlan) -> dict:
    return {
        "work_order_id": plan.work_order_id,
        "station_no": plan.station_no,
        "machine_id": plan.machine_id,
        "current_operation": plan.current_operation,
        "next_operation": plan.next_operation,
        "model_variant": plan.model_variant,
        "process_time": plan.process_time,
        "loading_unloading": plan.loading_unloading,
        "priority": plan.priority,
        "plan_type": plan.plan_type,
        "created_by": plan.created_by,
    }

def _compute_new_date(plan: ProductionPlan, data: RescheduleRequest) -> date:
    if data.mode == "next_week":
        return plan.plan_date + timedelta(days=7)
    if data.new_date:
        return data.new_date
    if data.days_offset is not None:
        return plan.plan_date + timedelta(days=data.days_offset)
    raise HTTPException(400, "Provide new_date, days_offset, or use mode=next_week")

def _validate_move_date(new_date: date) -> None:
    today = now_ist().date()
    if new_date < today:
        raise HTTPException(
            400,
            f"Cannot move plan to a past date ({new_date}). Earliest allowed date is {today}.",
        )

def _validate_may_complete_plan(plan: ProductionPlan) -> None:
    today = now_ist().date()
    if plan.plan_date > today:
        raise HTTPException(
            400,
            f"Cannot complete plan scheduled for {plan.plan_date}. "
            f"Completion is only allowed on or after the plan date (today: {today}).",
        )

def _validate_may_start_plan(plan: ProductionPlan) -> None:
    today = now_ist().date()
    if plan.plan_date > today:
        raise HTTPException(
            400,
            f"Cannot start plan scheduled for {plan.plan_date}. "
            f"Start is only allowed on or after the plan date (today: {today}).",
        )

def _should_split(plan: ProductionPlan, data: RescheduleRequest) -> bool:
    if data.mode == "split_remaining" or data.split_remaining:
        return plan.actual_qty > 0 and plan.actual_qty < plan.planned_qty
    return plan.status == "paused" and plan.actual_qty > 0 and plan.actual_qty < plan.planned_qty

def _apply_machine_relocation(plan: ProductionPlan, data: RescheduleRequest, db: Session) -> Optional[str]:
    if data.new_machine_id is None:
        return None
    if data.new_machine_id == plan.machine_id:
        return None
    machine = db.query(Machine).filter(Machine.id == data.new_machine_id).first()
    if not machine:
        raise HTTPException(404, "Target machine not found")
    if machine.status == "breakdown":
        raise HTTPException(400, f"Cannot move plan to {machine.name} — machine is in breakdown")
    old_m = db.query(Machine).filter(Machine.id == plan.machine_id).first() if plan.machine_id else None
    old_name = old_m.name if old_m else "unassigned"
    plan.machine_id = data.new_machine_id
    plan.station_no = data.new_station_no or machine.station_id
    return f" [Machine relocated: {old_name} → {machine.name}]"

def _mold_label(mold: ToolStock) -> str:
    bits = [mold.tool_code, mold.tool_name]
    if getattr(mold, "part_no", None):
        bits.append(str(mold.part_no))
    return " · ".join(b for b in bits if b)


def _prepare_plan_row(data: PlanCreate, db: Session, user) -> dict:
    """Resolve mold master fields into ProductionPlan column values."""
    mold = None
    if data.mold_id:
        mold = db.query(ToolStock).filter(ToolStock.id == data.mold_id, ToolStock.active == 1).first()
        if not mold:
            raise HTTPException(404, "Mold not found in mold master")
    elif data.work_order_id:
        wo = db.query(WorkOrder).filter(WorkOrder.id == data.work_order_id).first()
        if wo and getattr(wo, "mold_id", None):
            mold = db.query(ToolStock).filter(ToolStock.id == wo.mold_id).first()

    details = dict(data.mold_details or {})
    if mold:
        details.setdefault("mold_id", mold.id)
        details.setdefault("mold_code", mold.tool_code)
        details.setdefault("mold_name", mold.tool_name)
        details.setdefault("part_no", getattr(mold, "part_no", None))
        details.setdefault("standard_cavity", getattr(mold, "standard_cavity", None))
        details.setdefault("customer_name", getattr(mold, "customer_name", None))
        details.setdefault("rm_grade", getattr(mold, "rm_grade", None))
        details.setdefault("standard_cycle_time", float(mold.standard_cycle_time) if mold.standard_cycle_time is not None else None)
        details.setdefault("std_shot_weight", float(mold.std_shot_weight) if mold.std_shot_weight is not None else None)
        details.setdefault("std_part_weight", float(mold.std_part_weight) if mold.std_part_weight is not None else None)
        details.setdefault("std_runner_weight", float(mold.std_runner_weight) if mold.std_runner_weight is not None else None)
        details.setdefault("colour", getattr(mold, "colour", None))
        details.setdefault("target_per_shift", getattr(mold, "target_per_shift", None))
        details.setdefault("mold_status", getattr(mold, "tool_status", None))
        if "running_cavity" not in details or details.get("running_cavity") in (None, ""):
            details["running_cavity"] = getattr(mold, "running_cavity", None)
        if not details.get("cavity_name"):
            from ..tool_service import active_cavity
            cav = active_cavity(mold)
            if cav:
                details["cavity_name"] = cav.get("name")
                details["cavity_id"] = cav.get("id")
        if not details.get("assigned_by"):
            details["assigned_by"] = getattr(user, "username", None) or getattr(user, "full_name", None) or str(user.id)

    mold_label = _mold_label(mold) if mold else None
    cavity = (details.get("cavity_name") or "").strip() or "-"
    current_op = (data.current_operation or (mold.tool_name if mold else None) or mold_label or "IMM")[:50]
    next_op = (data.next_operation or cavity)[:50]
    model_variant = data.model_variant or mold_label
    process_time = data.process_time
    if process_time is None or process_time == "":
        process_time = details.get("standard_cycle_time")
    if process_time is None:
        raise HTTPException(400, "Standard cycle time is required (select a mold or enter process time)")
    loading = data.loading_unloading
    row = data.model_dump(exclude={"end_date", "shifts", "tool_shortage_ack", "mold_details"})
    row["mold_id"] = mold.id if mold else data.mold_id
    row["mold_details_json"] = json.dumps(details) if details else None
    row["current_operation"] = current_op
    row["next_operation"] = next_op
    row["model_variant"] = model_variant
    row["process_time"] = float(process_time)
    # CT = standard cycle time only (mold change time is not added on the planning form)
    row["loading_unloading"] = float(loading if loading is not None else 0)
    return row


@router.post("/")
async def create_plan(data: PlanCreate, db: Session = Depends(get_db),
                      user=Depends(get_current_user)):
    start = data.plan_date
    end   = data.end_date if data.end_date and data.end_date >= start else start
    delta = (end - start).days + 1
    shifts = _resolve_shifts(data)

    if data.work_order_id:
        wo = db.query(WorkOrder).filter(WorkOrder.id == data.work_order_id).first()
        if not wo:
            raise HTTPException(404, "Work order not found")
        if wo.status == "cancelled":
            raise HTTPException(400, "Cannot plan against a cancelled work order")
        existing_planned = db.query(func.coalesce(func.sum(ProductionPlan.planned_qty), 0)).filter(
            ProductionPlan.work_order_id == wo.id,
            ProductionPlan.status != "cancelled",
        ).scalar() or 0
        total_new = data.planned_qty * delta * len(shifts)
        if existing_planned + total_new > wo.target_qty:
            remaining = max(wo.target_qty - existing_planned, 0)
            slot_count = delta * len(shifts)
            max_per_slot = remaining // slot_count if slot_count else remaining
            raise HTTPException(
                400,
                f"Total planned qty ({total_new} pcs = {data.planned_qty} per shift × {slot_count} slots) "
                f"exceeds work order capacity ({remaining} pcs left to plan). "
                f"Max per shift for this date/shift range: {max_per_slot} pcs.",
            )

        # Tool stock / life forecast — notify planner; require ack if short / near EOL
        from ..tool_service import build_forecast, log_event, find_tool
        forecast = build_forecast(db, work_order_id=data.work_order_id, planned_qty=total_new)
        if forecast.get("blocks_plan"):
            raise HTTPException(status_code=409, detail={
                "code": "tool_eol_blocked",
                "message": forecast.get("message"),
                "forecast": forecast,
            })
        if forecast.get("requires_ack") and not data.tool_shortage_ack:
            raise HTTPException(status_code=409, detail={
                "code": "tool_forecast_ack_required",
                "message": forecast.get("message"),
                "forecast": forecast,
            })
        if data.tool_shortage_ack and forecast.get("requires_ack"):
            from ..models import ToolStock
            for trow in forecast.get("tools") or []:
                if not trow.get("requires_ack"):
                    continue
                tool = None
                if trow.get("tool_id"):
                    tool = db.query(ToolStock).filter(ToolStock.id == trow["tool_id"]).first()
                if not tool:
                    tool = find_tool(db, tool_code=trow.get("tool_code"), tool_name=trow.get("tool_name"))
                if tool:
                    log_event(
                        db, tool, "forecast_ack",
                        user_id=user.id,
                        work_order_id=data.work_order_id,
                        notes=f"Planner acknowledged forecast: {trow.get('message')}",
                        acknowledged_by=user.id,
                    )

    created = []
    base = _prepare_plan_row(data, db, user)
    for i in range(delta):
        for shift in shifts:
            day_data = {**base, "plan_date": start + timedelta(days=i), "shift": shift}
            plan = ProductionPlan(**day_data, created_by=user.id, created_at=now_ist())
            db.add(plan)
            db.flush()
            created.append(plan)
    if data.work_order_id:
        sync_work_order_after_plan_change(db, data.work_order_id)
    db.commit()
    for plan in created:
        db.refresh(plan)
        await manager.broadcast({"type": "plan_created", "plan_id": plan.id,
                                  "station_no": plan.station_no, "shift": plan.shift,
                                  "current_operation": plan.current_operation, "next_operation": plan.next_operation,
                                  "planned_qty": plan.planned_qty})
    return created if delta > 1 else created[0]

@router.get("/")
def get_plans(
    plan_date: Optional[date] = None,
    shift: Optional[str] = None,
    station_no: Optional[int] = None,
    week: Optional[int] = None,
    week_start: Optional[date] = None,
    week_end: Optional[date] = None,
    month: Optional[int] = None,
    year: Optional[int] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db),
    _=Depends(get_current_user)
):
    q = db.query(ProductionPlan)
    if plan_date:  q = q.filter(ProductionPlan.plan_date == plan_date)
    if shift:      q = q.filter(ProductionPlan.shift == shift)
    if station_no:    q = q.filter(ProductionPlan.station_no == station_no)
    if month:      q = q.filter(extract("month", ProductionPlan.plan_date) == month)
    if year:       q = q.filter(extract("year",  ProductionPlan.plan_date) == year)
    if week:       q = q.filter(extract("week",  ProductionPlan.plan_date) == week)
    if week_start: q = q.filter(ProductionPlan.plan_date >= week_start)
    if week_end:   q = q.filter(ProductionPlan.plan_date <= week_end)
    if date_from:  q = q.filter(ProductionPlan.plan_date >= date_from)
    if date_to:    q = q.filter(ProductionPlan.plan_date <= date_to)
    if status:     q = q.filter(ProductionPlan.status == status)
    plans = q.order_by(ProductionPlan.plan_date, ProductionPlan.shift, ProductionPlan.priority).all()

    # Non-blocking: refresh actuals in background (GET must stay fast)
    try:
        from ..plan_actuals_sync import schedule_plan_actuals_sync
        sk: dict = {}
        if plan_date:
            sk["entry_date"] = plan_date
        else:
            if date_from or week_start:
                sk["date_from"] = date_from or week_start
            if date_to or week_end:
                sk["date_to"] = date_to or week_end
        if shift:
            sk["shift"] = shift
        schedule_plan_actuals_sync(**sk)
    except Exception as exc:
        print(f"[Plans] background actual sync schedule skipped: {exc}")

    return [_plan_dict(p, db) for p in plans]

@router.get("/summary")
def get_summary(
    plan_date: Optional[date] = None,
    shift: Optional[str] = None,
    station_no: Optional[int] = None,
    month: Optional[int] = None,
    year: Optional[int] = None,
    week_start: Optional[date] = None,
    week_end: Optional[date] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    _=Depends(get_current_user)
):
    q = db.query(ProductionPlan)
    if plan_date:  q = q.filter(ProductionPlan.plan_date == plan_date)
    if shift:      q = q.filter(ProductionPlan.shift == shift)
    if station_no:    q = q.filter(ProductionPlan.station_no == station_no)
    if month:      q = q.filter(extract("month", ProductionPlan.plan_date) == month)
    if year:       q = q.filter(extract("year",  ProductionPlan.plan_date) == year)
    if week_start: q = q.filter(ProductionPlan.plan_date >= week_start)
    if week_end:   q = q.filter(ProductionPlan.plan_date <= week_end)
    if date_from:  q = q.filter(ProductionPlan.plan_date >= date_from)
    if date_to:    q = q.filter(ProductionPlan.plan_date <= date_to)

    # Non-blocking background refresh — summary uses persisted actual_qty
    try:
        from ..plan_actuals_sync import schedule_plan_actuals_sync
        sk: dict = {}
        if plan_date:
            sk["entry_date"] = plan_date
        else:
            if date_from or week_start:
                sk["date_from"] = date_from or week_start
            if date_to or week_end:
                sk["date_to"] = date_to or week_end
        if shift:
            sk["shift"] = shift
        schedule_plan_actuals_sync(**sk)
    except Exception as exc:
        print(f"[Plans] summary background sync schedule skipped: {exc}")

    plans = q.all()
    total_planned = sum(p.planned_qty for p in plans)
    total_actual  = sum(p.actual_qty  for p in plans)
    return {
        "total_plans": len(plans),
        "total_planned": total_planned,
        "total_actual": total_actual,
        "achievement_pct": round(total_actual / total_planned * 100, 1) if total_planned else 0,
        "by_status": {s: sum(1 for p in plans if p.status == s)
                      for s in ["pending","running","completed","paused","cancelled","aborted","incomplete"]},
        "by_shift": {sh: {"planned": sum(p.planned_qty for p in plans if p.shift == sh),
                           "actual":  sum(p.actual_qty  for p in plans if p.shift == sh)}
                     for sh in ["A","B"]}
    }

@router.get("/{plan_id}")
def get_plan(plan_id: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    """Fetch a single plan by ID. Returns the plan even if its linked work order was deleted."""
    plan = db.query(ProductionPlan).filter(ProductionPlan.id == plan_id).first()
    if not plan:
        raise HTTPException(404, "Plan not found")
    return _plan_dict(plan, db)


def _plan_dict(plan: ProductionPlan, db: Session) -> dict:
    """Serialize plan with linked model-change interlock state."""
    d = {c.name: getattr(plan, c.name) for c in plan.__table__.columns}
    for key in ("plan_date", "created_at", "updated_at"):
        val = d.get(key)
        if val is not None and hasattr(val, "isoformat"):
            d[key] = val.isoformat()
    for key in ("process_time", "loading_unloading"):
        if d.get(key) is not None:
            d[key] = float(d[key])
    raw_details = d.pop("mold_details_json", None) or getattr(plan, "mold_details_json", None)
    mold_details = {}
    if raw_details:
        try:
            mold_details = json.loads(raw_details) if isinstance(raw_details, str) else (raw_details or {})
        except (json.JSONDecodeError, TypeError):
            mold_details = {}
    d["mold_details"] = mold_details
    if getattr(plan, "mold_id", None):
        mold = db.query(ToolStock).filter(ToolStock.id == plan.mold_id).first()
        if mold:
            d["mold_code"] = mold.tool_code
            d["mold_name"] = mold.tool_name
            d["mold_label"] = _mold_label(mold)
    pending = (
        db.query(ModelChangeRequest)
        .filter(
            ModelChangeRequest.plan_id == plan.id,
            ModelChangeRequest.status == "pending",
        )
        .order_by(ModelChangeRequest.id.desc())
        .first()
    )
    active = (
        db.query(ModelChangeRequest)
        .filter(
            ModelChangeRequest.plan_id == plan.id,
            ModelChangeRequest.status.in_(["approved", "in_progress"]),
        )
        .order_by(ModelChangeRequest.id.desc())
        .first()
    )
    mcr = pending or active
    d["model_change_request_id"] = mcr.id if mcr else None
    d["model_change_status"] = mcr.status if mcr else None
    d["awaiting_model_change"] = bool(pending)
    return d


def _same_part_in_previous_shift(
    db: Session,
    machine_id: int,
    model_variant: str,
    plan_date: date,
    shift_id: str,
) -> bool:
    """Check if the same model_variant was running/completed on this machine
    in the immediately preceding shift (including cross-day C→A)."""
    from .hourly_output import _load_config, _previous_shift_info

    cfg = _load_config(db)
    prev_shift_id, prev_date = _previous_shift_info(cfg, shift_id, plan_date)
    if not prev_shift_id or not prev_date:
        return False

    prev_plan = (
        db.query(ProductionPlan)
        .filter(
            ProductionPlan.machine_id == machine_id,
            ProductionPlan.plan_date == prev_date,
            ProductionPlan.shift == prev_shift_id,
            ProductionPlan.model_variant == model_variant,
            ProductionPlan.status.in_(["running", "completed", "paused"]),
        )
        .first()
    )
    return prev_plan is not None


def _previous_model_on_machine(db: Session, machine_id: int, exclude_plan_id: int) -> str:
    prev = (
        db.query(ProductionPlan)
        .filter(
            ProductionPlan.machine_id == machine_id,
            ProductionPlan.id != exclude_plan_id,
            ProductionPlan.status.in_(["running", "completed", "paused"]),
        )
        .order_by(ProductionPlan.updated_at.desc(), ProductionPlan.id.desc())
        .first()
    )
    if prev and prev.model_variant:
        return prev.model_variant
    return "—"


@router.patch("/{plan_id}/status")
async def update_status(plan_id: int, data: PlanUpdate, db: Session = Depends(get_db),
                        user=Depends(get_current_user)):
    plan = db.query(ProductionPlan).filter(ProductionPlan.id == plan_id).first()
    if not plan:
        raise HTTPException(404, "Plan not found")
    if data.status == "completed":
        _validate_may_complete_plan(plan)
    if data.status == "running" and plan.status != "running":
        _validate_may_start_plan(plan)

    if data.status == "aborted":
        if plan.status not in ("paused", "running"):
            raise HTTPException(400, "Only running or paused plans can be aborted")
        plan.status = "aborted"
        plan.updated_at = now_ist()
        machine = db.query(Machine).filter(Machine.id == plan.machine_id).first() if plan.machine_id else None
        if machine and machine.status == "running":
            other_running = (
                db.query(ProductionPlan)
                .filter(
                    ProductionPlan.machine_id == plan.machine_id,
                    ProductionPlan.id != plan.id,
                    ProductionPlan.status == "running",
                )
                .first()
            )
            if not other_running:
                machine.status = "idle"
        from .work_orders import sync_work_order_after_plan_change
        sync_work_order_after_plan_change(db, plan.work_order_id)
        db.commit()
        await manager.broadcast({"type": "plan_updated", "plan_id": plan_id, "status": "aborted"})
        return _plan_dict(plan, db)

    # ── Interlock: pending → running requires model-change approval ──
    if data.status == "running" and plan.status == "pending":
        if not plan.machine_id:
            machine = (
                db.query(Machine)
                .filter(Machine.station_id == plan.station_no)
                .order_by(Machine.id)
                .first()
            )
            if not machine:
                raise HTTPException(
                    400,
                    "No machine found for this station — assign a machine on the plan before starting",
                )
            plan.machine_id = machine.id

        # Skip model-change interlock when the same part was already running
        # on this machine in the previous shift (no actual model change needed)
        same_part_continues = False
        if plan.machine_id and plan.model_variant:
            same_part_continues = _same_part_in_previous_shift(
                db, plan.machine_id, plan.model_variant,
                plan.plan_date, plan.shift,
            )

        if same_part_continues:
            if plan.plan_type != "trial" and plan.machine_id:
                conflicting = db.query(ProductionPlan).filter(
                    ProductionPlan.machine_id == plan.machine_id,
                    ProductionPlan.id != plan.id,
                    ProductionPlan.status == "running",
                ).all()
                for conflict in conflicting:
                    conflict.status = "paused"
                    conflict.updated_at = now_ist()
                    await manager.broadcast({"type": "plan_updated", "plan_id": conflict.id, "status": "paused"})

            plan.status = "running"
            plan.updated_at = now_ist()
            machine = db.query(Machine).filter(Machine.id == plan.machine_id).first()
            if machine and machine.status not in ("breakdown", "offline", "setting_change", "alarm"):
                machine.status = "running"
            db.commit()
            await manager.broadcast({
                "type": "plan_started", "plan_id": plan_id,
                "machine_id": plan.machine_id, "station_no": plan.station_no,
                "current_operation": plan.current_operation, "next_operation": plan.next_operation,
                "model_variant": plan.model_variant,
                "process_time": float(plan.process_time) if plan.process_time is not None else None,
                "loading_unloading": float(plan.loading_unloading) if plan.loading_unloading is not None else None,
                "shift": plan.shift, "plan_date": str(plan.plan_date),
            })
            return {
                **_plan_dict(plan, db),
                "message": "Same part continues from previous shift — started without model change",
            }

        existing_pending = (
            db.query(ModelChangeRequest)
            .filter(
                ModelChangeRequest.plan_id == plan.id,
                ModelChangeRequest.status == "pending",
            )
            .first()
        )
        if existing_pending:
            # Idempotent: Start already raised a request — do not error on repeat clicks
            return {
                **_plan_dict(plan, db),
                "model_change_pending": True,
                "model_change_request_id": existing_pending.id,
                "message": (
                    f"Model change request #{existing_pending.id} is awaiting approval "
                    "on the Model Change page"
                ),
            }

        existing_active = (
            db.query(ModelChangeRequest)
            .filter(
                ModelChangeRequest.plan_id == plan.id,
                ModelChangeRequest.status.in_(["approved", "in_progress"]),
            )
            .first()
        )
        if existing_active:
            # Approval already applied plan start — treat as already running
            # Auto-pause any other running plan on the same machine (unless trial)
            if plan.plan_type != "trial" and plan.machine_id:
                conflicting = db.query(ProductionPlan).filter(
                    ProductionPlan.machine_id == plan.machine_id,
                    ProductionPlan.id != plan.id,
                    ProductionPlan.status == "running",
                ).all()
                for conflict in conflicting:
                    conflict.status = "paused"
                    conflict.updated_at = now_ist()
            plan.status = "running"
            plan.updated_at = now_ist()
            db.commit()
            return {
                **_plan_dict(plan, db),
                "message": "Plan is running (model change already approved)",
            }

        from_model = _previous_model_on_machine(db, plan.machine_id, plan.id)
        to_model = (plan.model_variant or plan.current_operation or "—").strip() or "—"
        mcr = ModelChangeRequest(
            machine_id=plan.machine_id,
            plan_id=plan.id,
            requested_by=user.id,
            from_model=from_model,
            to_model=to_model,
            ideal_minutes=60,
            shift=plan.shift or "A",
            entry_date=plan.plan_date or now_ist().date(),
            reason="setting_change",
            status="pending",
            created_at=now_ist(),
        )
        db.add(mcr)
        # Keep plan pending until supervisor approves on Model Change page
        plan.updated_at = now_ist()
        db.commit()
        db.refresh(mcr)
        await manager.broadcast({
            "type": "model_change_request",
            "id": mcr.id,
            "machine_id": plan.machine_id,
            "plan_id": plan.id,
            "from_model": from_model,
            "to_model": to_model,
            "status": "pending",
        })
        await manager.broadcast({"type": "plan_updated", "plan_id": plan_id, "status": plan.status})
        return {
            **_plan_dict(plan, db),
            "message": (
                f"Model change request #{mcr.id} raised ({from_model} → {to_model}). "
                "Approve on Model Change page to start the plan and apply the part on WI."
            ),
            "model_change_pending": True,
        }

    if data.status:
        plan.status = data.status
    if data.planned_qty is not None:
        plan.planned_qty = data.planned_qty
    if data.priority is not None:
        plan.priority = data.priority
    if data.plan_type:
        plan.plan_type = data.plan_type
    if data.notes is not None:
        plan.notes = data.notes
    plan.updated_at = now_ist()

    # Resume from paused (no new model-change gate)
    if data.status == "running":
        # Auto-pause any other running plan on the same machine, unless this is a trial plan
        # (trial = intentional concurrent run, e.g. setup verification alongside production)
        if plan.plan_type != "trial" and plan.machine_id:
            conflicting = db.query(ProductionPlan).filter(
                ProductionPlan.machine_id == plan.machine_id,
                ProductionPlan.id != plan.id,
                ProductionPlan.status == "running",
            ).all()
            for conflict in conflicting:
                conflict.status = "paused"
                conflict.updated_at = now_ist()
                await manager.broadcast({"type": "plan_updated", "plan_id": conflict.id, "status": "paused"})

        machine = db.query(Machine).filter(Machine.id == plan.machine_id).first()
        if machine and machine.status not in ("breakdown", "offline", "setting_change", "alarm"):
            machine.status = "running"
        db.commit()
        await manager.broadcast({
            "type": "plan_started", "plan_id": plan_id,
            "machine_id": plan.machine_id, "station_no": plan.station_no,
            "current_operation": plan.current_operation, "next_operation": plan.next_operation,
            "model_variant": plan.model_variant,
            "process_time": float(plan.process_time) if plan.process_time is not None else None,
            "loading_unloading": float(plan.loading_unloading) if plan.loading_unloading is not None else None,
            "shift": plan.shift, "plan_date": str(plan.plan_date),
        })
    elif data.status == "completed":
        machine = db.query(Machine).filter(Machine.id == plan.machine_id).first()
        if machine:
            machine.status = "idle"
        sync_work_order_after_plan_change(db, plan.work_order_id)
        db.commit()
        await manager.broadcast({"type": "plan_completed", "plan_id": plan_id,
                                  "machine_id": plan.machine_id, "station_no": plan.station_no})
    else:
        db.commit()
        await manager.broadcast({"type": "plan_updated", "plan_id": plan_id, "status": data.status})
    return _plan_dict(plan, db)

@router.patch("/{plan_id}/actual")
async def update_actual_qty(plan_id: int, data: ActualQtyUpdate, db: Session = Depends(get_db),
                             user=Depends(get_current_user)):
    """Called by manual entry, MQTT bridge, Modbus bridge, or OPC-UA bridge"""
    plan = db.query(ProductionPlan).filter(ProductionPlan.id == plan_id).first()
    if not plan: raise HTTPException(404, "Plan not found")
    prev_actual = int(plan.actual_qty or 0)
    plan.actual_qty = data.actual_qty
    plan.updated_at = now_ist()
    if data.actual_qty >= plan.planned_qty and plan.status == "running":
        _validate_may_complete_plan(plan)
        plan.status = "completed"
        machine = db.query(Machine).filter(Machine.id == plan.machine_id).first()
        if machine: machine.status = "idle"
    try:
        from ..tool_service import apply_consumption_safe
        apply_consumption_safe(db, plan, prev_actual, data.actual_qty, user_id=getattr(user, "id", None))
    except Exception as exc:
        print(f"[WARN] tool consumption skipped: {exc}")
    sync_work_order_after_plan_change(db, plan.work_order_id)
    db.commit()
    await manager.broadcast({
        "type": "actual_qty_updated", "plan_id": plan_id,
        "actual_qty": data.actual_qty, "planned_qty": plan.planned_qty,
        "source": data.source, "station_no": plan.station_no
    })
    return plan

@router.delete("/{plan_id}")
async def delete_plan(plan_id: int, db: Session = Depends(get_db),
                      user=Depends(require_role("supervisor", "admin"))):
    plan = db.query(ProductionPlan).filter(ProductionPlan.id == plan_id).first()
    if not plan: raise HTTPException(404, "Plan not found")
    db.delete(plan)
    db.commit()
    await manager.broadcast({"type": "plan_deleted", "plan_id": plan_id})
    return {"ok": True}


@router.post("/{plan_id}/reschedule")
async def reschedule_plan(
    plan_id: int,
    data: RescheduleRequest,
    db: Session = Depends(get_db),
    user=Depends(require_role("supervisor", "admin")),
):
    plan = db.query(ProductionPlan).filter(ProductionPlan.id == plan_id).first()
    if not plan:
        raise HTTPException(404, "Plan not found")
    if plan.status in ("completed", "cancelled"):
        raise HTTPException(400, f"Cannot reschedule a {plan.status} plan")
    if plan.status == "running":
        raise HTTPException(400, "Pause the plan before moving it")

    new_date = _compute_new_date(plan, data)
    _validate_move_date(new_date)
    new_shift = data.new_shift or plan.shift
    old_date = plan.plan_date
    created = None

    if _should_split(plan, data):
        remaining = plan.planned_qty - plan.actual_qty
        _validate_may_complete_plan(plan)
        plan.planned_qty = plan.actual_qty
        plan.status = "completed"
        suffix = f" [Partial: {remaining} pcs moved to {new_date} shift {new_shift}]"
        plan.notes = ((plan.notes or "") + suffix).strip()
        plan.updated_at = now_ist()

        created = ProductionPlan(
            **_plan_copy_fields(plan),
            plan_date=new_date,
            shift=new_shift,
            planned_qty=remaining,
            actual_qty=0,
            status="pending",
            notes=f"Moved from plan #{plan_id} ({old_date} shift {plan.shift})",
            created_at=now_ist(),
            updated_at=now_ist(),
        )
        machine_note = _apply_machine_relocation(created, data, db)
        if machine_note:
            created.notes = (created.notes + machine_note).strip()
        db.add(created)
    else:
        old_shift = plan.shift
        plan.plan_date = new_date
        plan.shift = new_shift
        if plan.status == "paused":
            plan.status = "pending"
        suffix = f" [Rescheduled from {old_date} shift {old_shift}]"
        plan.notes = ((plan.notes or "") + suffix).strip()
        machine_note = _apply_machine_relocation(plan, data, db)
        if machine_note:
            plan.notes = (plan.notes + machine_note).strip()
        plan.updated_at = now_ist()

    if plan.work_order_id:
        sync_work_order_after_plan_change(db, plan.work_order_id)
    db.commit()
    db.refresh(plan)
    if created:
        db.refresh(created)

    await manager.broadcast({
        "type": "plan_rescheduled",
        "plan_id": plan_id,
        "new_plan_id": created.id if created else None,
        "new_date": str(new_date),
    })
    return {
        "original": plan,
        "new_plan": created,
        "split": created is not None,
    }


@router.post("/bulk-reschedule")
async def bulk_reschedule(
    data: BulkRescheduleRequest,
    db: Session = Depends(get_db),
    user=Depends(require_role("supervisor", "admin")),
):
    if not data.plan_ids:
        raise HTTPException(400, "No plans selected")

    plans = db.query(ProductionPlan).filter(ProductionPlan.id.in_(data.plan_ids)).all()
    if len(plans) != len(data.plan_ids):
        raise HTTPException(404, "One or more plans not found")

    movable = [p for p in plans if p.status in ("pending", "paused")]
    if not movable:
        raise HTTPException(400, "No pending or paused plans to move")

    if data.new_start_date:
        anchor = min(p.plan_date for p in movable)
        offset_days = (data.new_start_date - anchor).days
    elif data.mode == "next_week":
        offset_days = data.days_offset if data.days_offset else 7
    else:
        offset_days = data.days_offset or 0

    today = now_ist().date()
    if data.new_start_date and data.new_start_date < today:
        raise HTTPException(
            400,
            f"Start date cannot be in the past ({data.new_start_date}). Earliest allowed date is {today}.",
        )

    invalid_moves = []
    for plan in movable:
        candidate = plan.plan_date + timedelta(days=offset_days)
        if candidate < today:
            invalid_moves.append(plan.id)
    if invalid_moves:
        raise HTTPException(
            400,
            f"Cannot move plan(s) {invalid_moves} to a past date. Earliest allowed date is {today}.",
        )

    results = []
    work_order_ids = set()

    for plan in movable:
        if plan.status == "running":
            continue
        req = RescheduleRequest(
            mode="custom",
            days_offset=offset_days,
            split_remaining=data.split_remaining,
        )
        new_date = plan.plan_date + timedelta(days=offset_days)
        new_shift = plan.shift
        old_date = plan.plan_date
        created = None

        if _should_split(plan, req):
            remaining = plan.planned_qty - plan.actual_qty
            _validate_may_complete_plan(plan)
            plan.planned_qty = plan.actual_qty
            plan.status = "completed"
            plan.notes = ((plan.notes or "") + f" [Bulk move: {remaining} pcs → {new_date}]").strip()
            plan.updated_at = now_ist()
            created = ProductionPlan(
                **_plan_copy_fields(plan),
                plan_date=new_date,
                shift=new_shift,
                planned_qty=remaining,
                actual_qty=0,
                status="pending",
                notes=f"Bulk moved from plan #{plan.id} ({old_date})",
                created_at=now_ist(),
                updated_at=now_ist(),
            )
            db.add(created)
        else:
            plan.plan_date = new_date
            if plan.status == "paused":
                plan.status = "pending"
            plan.notes = ((plan.notes or "") + f" [Bulk moved from {old_date}]").strip()
            plan.updated_at = now_ist()

        if plan.work_order_id:
            work_order_ids.add(plan.work_order_id)
        results.append({"plan_id": plan.id, "new_date": str(new_date), "split": created is not None})

    for wo_id in work_order_ids:
        sync_work_order_after_plan_change(db, wo_id)
    db.commit()

    await manager.broadcast({"type": "plans_bulk_rescheduled", "count": len(results)})
    return {"moved": len(results), "results": results, "days_offset": offset_days}

@router.get("/pipeline/{station_no}")
def get_pipeline(station_no: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    """Returns ordered queue for a machine — next part to load is first"""
    plans = db.query(ProductionPlan).filter(
        ProductionPlan.station_no == station_no,
        ProductionPlan.status.in_(["pending", "running"])
    ).order_by(ProductionPlan.plan_date, ProductionPlan.shift, ProductionPlan.priority).all()
    return [_plan_dict(p, db) for p in plans]


def _style_header_row(ws, ncols):
    hdr_fill = PatternFill("solid", fgColor="1E3A5F")
    hdr_font = Font(bold=True, color="FFFFFF")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

def _build_excel(plans, work_orders: Optional[dict] = None) -> io.BytesIO:
    work_orders = work_orders or {}
    wb = openpyxl.Workbook()
    now = datetime.now()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    total_planned = sum(p.planned_qty for p in plans)
    total_actual  = sum(p.actual_qty  for p in plans)
    achievement   = round(total_actual / total_planned * 100, 1) if total_planned else 0
    by_status = {s: sum(1 for p in plans if p.status == s)
                 for s in ["pending", "running", "completed", "paused", "cancelled"]}
    summary_rows = [
        ["PRODUCTION PLANNING REPORT"],
        [""],
        ["Report Date:", now.strftime("%d-%m-%Y")],
        ["Period:", f"{plans[0].plan_date} to {plans[-1].plan_date}" if len(plans) > 1 else (str(plans[0].plan_date) if plans else "")],
        [""],
        ["SUMMARY METRICS"],
        ["Total Plans:",    len(plans)],
        ["Planned Quantity:", total_planned],
        ["Actual Quantity:",  total_actual],
        ["Achievement %:",    str(achievement) + "%"],
        [""],
        ["STATUS BREAKDOWN"],
        ["Pending:",   by_status["pending"]],
        ["Running:",   by_status["running"]],
        ["Completed:", by_status["completed"]],
        ["Paused:",    by_status["paused"]],
        ["Cancelled:", by_status["cancelled"]],
    ]
    for row in summary_rows:
        ws1.append(row)
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 18

    # ── Sheet 2: Plans Details ────────────────────────────────────────────────
    ws2 = wb.create_sheet("Plans Details")
    headers2 = ["Date","Shift","Station","Work Order No","WO Target Qty","Current Operation","Next Operation","Process Time (s)",
                "Loading/Unloading (s)","Cycle Time (s)","Type","Priority",
                "Planned Qty","Actual Qty","Achievement %","Status","Notes"]
    ws2.append(headers2)
    _style_header_row(ws2, len(headers2))
    for p in plans:
        pct = round(p.actual_qty / p.planned_qty * 100, 1) if p.planned_qty else 0
        wo = work_orders.get(p.work_order_id) if p.work_order_id else None
        ws2.append([str(p.plan_date), p.shift, getattr(p, '_station_label', p.station_no),
                    wo.work_order_no if wo else "", wo.target_qty if wo else "",
                    p.current_operation, p.next_operation,
                    p.process_time, p.loading_unloading, p.process_time + p.loading_unloading,
                    p.plan_type, p.priority, p.planned_qty, p.actual_qty,
                    str(pct) + "%", p.status, p.notes or ""])
    col_widths2 = [12,8,6,14,12,14,14,16,20,14,12,10,12,12,14,12,20]
    for i, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w

    # ── Sheet 3: By Pair ──────────────────────────────────────────────────────
    ws3 = wb.create_sheet("By Station")
    headers3 = ["Station No","Total Plans","Planned Qty","Actual Qty",
                "Achievement %","Completed","Running","Pending"]
    ws3.append(headers3)
    _style_header_row(ws3, len(headers3))
    station_stats = {}
    for p in plans:
        key = getattr(p, '_station_label', str(p.station_no))
        ps = station_stats.setdefault(key, {"planned": 0, "actual": 0, "total": 0,
                                         "completed": 0, "running": 0, "pending": 0})
        ps["total"] += 1; ps["planned"] += p.planned_qty; ps["actual"] += p.actual_qty
        if p.status in ps: ps[p.status] += 1
    for pn, ps in sorted(station_stats.items(), key=lambda x: str(x[0])):
        pct = round(ps["actual"] / ps["planned"] * 100, 1) if ps["planned"] else 0
        ws3.append([pn, ps["total"], ps["planned"], ps["actual"],
                    str(pct) + "%", ps["completed"], ps["running"], ps["pending"]])
    for i, w in enumerate([10,12,12,12,14,12,12,12], 1):
        ws3.column_dimensions[ws3.cell(row=1, column=i).column_letter].width = w

    # ── Sheet 4: By Shift ─────────────────────────────────────────────────────
    ws4 = wb.create_sheet("By Shift")
    headers4 = ["Shift","Total Plans","Planned Qty","Actual Qty",
                "Achievement %","Completed","Running"]
    ws4.append(headers4)
    _style_header_row(ws4, len(headers4))
    shift_stats = {}
    for p in plans:
        ss = shift_stats.setdefault(p.shift, {"planned": 0, "actual": 0, "total": 0,
                                              "completed": 0, "running": 0})
        ss["total"] += 1; ss["planned"] += p.planned_qty; ss["actual"] += p.actual_qty
        if p.status in ss: ss[p.status] += 1
    for sh, ss in sorted(shift_stats.items()):
        pct = round(ss["actual"] / ss["planned"] * 100, 1) if ss["planned"] else 0
        ws4.append([sh, ss["total"], ss["planned"], ss["actual"],
                    str(pct) + "%", ss["completed"], ss["running"]])
    for i, w in enumerate([10,12,12,12,14,12,12], 1):
        ws4.column_dimensions[ws4.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ExportParams(BaseModel):
    plan_date: Optional[str] = None
    date_from: Optional[str] = None
    date_to: Optional[str] = None
    month: Optional[int] = None
    year: Optional[int] = None
    shift: Optional[str] = None
    station_no: Optional[int] = None


class EmailReportRequest(ExportParams):
    recipients: List[str]
    subject: Optional[str] = "Production Planning Report"


def _filter_plans(q, params: ExportParams):
    if params.plan_date: q = q.filter(ProductionPlan.plan_date == params.plan_date)
    if params.date_from: q = q.filter(ProductionPlan.plan_date >= params.date_from)
    if params.date_to:   q = q.filter(ProductionPlan.plan_date <= params.date_to)
    if params.shift:     q = q.filter(ProductionPlan.shift == params.shift)
    if params.station_no:   q = q.filter(ProductionPlan.station_no == params.station_no)
    if params.month:     q = q.filter(extract("month", ProductionPlan.plan_date) == params.month)
    if params.year:      q = q.filter(extract("year",  ProductionPlan.plan_date) == params.year)
    return q


def _work_orders_for_plans(db: Session, plans) -> dict:
    wo_ids = {p.work_order_id for p in plans if p.work_order_id}
    if not wo_ids:
        return {}
    rows = db.query(WorkOrder).filter(WorkOrder.id.in_(wo_ids)).all()
    return {w.id: w for w in rows}


@router.post("/export")
def export_excel(params: ExportParams, db: Session = Depends(get_db),
                 _=Depends(get_current_user)):
    q = _filter_plans(db.query(ProductionPlan), params)
    plans = q.order_by(ProductionPlan.plan_date, ProductionPlan.shift, ProductionPlan.priority).all()
    buf = _build_excel(plans, _work_orders_for_plans(db, plans))
    filename = f"production_plan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/email-logs")
def get_email_logs(db: Session = Depends(get_db), _=Depends(get_current_user)):
    logs = db.query(EmailLog).order_by(EmailLog.sent_at.desc()).limit(100).all()
    return [{"id": l.id, "sent_at": l.sent_at, "recipients": l.recipients,
             "subject": l.subject, "report_type": l.report_type,
             "status": l.status, "error_msg": l.error_msg} for l in logs]


@router.post("/email-report")
def email_report(req: EmailReportRequest, db: Session = Depends(get_db),
                 user=Depends(get_current_user)):
    q = _filter_plans(db.query(ProductionPlan), req)
    plans = q.order_by(ProductionPlan.plan_date, ProductionPlan.shift, ProductionPlan.priority).all()
    buf = _build_excel(plans, _work_orders_for_plans(db, plans))
    filename = f"production_plan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    from ..models import EmailSmtpConfig
    cfg = db.query(EmailSmtpConfig).first()
    smtp_host = cfg.smtp_server if cfg else os.getenv("SMTP_HOST", "smtp.gmail.com")
    smtp_port = cfg.smtp_port   if cfg else int(os.getenv("SMTP_PORT", 587))
    smtp_user = cfg.email_address  if cfg else os.getenv("SMTP_USER", "")
    smtp_pass = cfg.email_password if cfg else os.getenv("SMTP_PASS", "")

    if not smtp_user or not smtp_pass:
        raise HTTPException(400, "SMTP credentials not configured")

    msg = MIMEMultipart()
    msg["From"]    = smtp_user
    msg["To"]      = ", ".join(req.recipients)
    msg["Subject"] = req.subject
    msg.attach(MIMEText(f"Please find attached the production planning report generated on {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}.", "plain"))
    part = MIMEBase("application", "octet-stream")
    part.set_payload(buf.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f"attachment; filename={filename}")
    msg.attach(part)

    log = EmailLog(sent_at=now_ist(), recipients=", ".join(req.recipients),
                   subject=req.subject, report_type="planning",
                   sent_by=user.id if user else None)
    try:
        with smtplib.SMTP(smtp_host, smtp_port) as server:
            server.starttls()
            server.login(smtp_user, smtp_pass)
            server.sendmail(smtp_user, req.recipients, msg.as_string())
        log.status = "sent"
    except Exception as e:
        log.status = "failed"
        log.error_msg = str(e)
        db.add(log)
        db.commit()
        raise HTTPException(500, f"Email send failed: {str(e)}")
    db.add(log)
    db.commit()
    return {"ok": True, "sent_to": req.recipients}
